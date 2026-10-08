"""Build the demo topology workbook (``data/topology.xlsx``) from Hans's ``demo.xlsx``.

The hand-made ``demo.xlsx`` describes the demo network (2 route reflectors, 8 core
routers, 11 PE sites in half-open rings, one access switch per PE site) but uses
placeholder device models with Cisco-style port names. Firestarter renders Nokia
SR OS, so this script rewrites the workbook onto real SR OS models (whose port
recipes live in ``DeviceFactory``), fills in fictional site addresses, adds the
prefix / resource pools the demo services need, and drops the ``report_*`` tabs
(those are outputs the pipeline writes back).

Usage::

    python scripts/build_demo_workbook.py <path-to-demo.xlsx> data/topology.xlsx
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

MODEL_BY_ROLE = {
    "rr": "7250-IXR-X3",
    "core": "7750-SR-1x-48d",
    "pe": "7250-IXR-e2-100",
}

CITIES = [
    "Amsterdam",
    "Rotterdam",
    "Utrecht",
    "Eindhoven",
    "Groningen",
    "Tilburg",
    "Almere",
    "Breda",
    "Nijmegen",
    "Apeldoorn",
    "Haarlem",
    "Arnhem",
    "Enschede",
    "Amersfoort",
    "Zaanstad",
    "Den Bosch",
    "Haarlemmermeer",
    "Zwolle",
    "Zoetermeer",
    "Leiden",
]

PREFIX_POOLS = [
    # name, type, prefix
    ("core_loopback0_pool_demo", "loopback", "10.0.0.0/24"),
    ("core_loopback1_pool_demo", "loopback", "10.0.2.0/24"),
    ("pe_loopback0_pool_demo", "loopback", "10.0.1.0/24"),
    ("pe_loopback1_pool_demo", "loopback", "10.0.3.0/24"),
    ("p2p_pool_demo", "p2p", "10.0.4.0/22"),
    ("delegated_pool_ce_mgmt_demo", "delegated", "10.0.8.0/22"),
]

RESOURCE_POOLS = [
    ("evpn_esi_pool_demo", 1000, 2000),
    ("sr_core_pool_demo", 16000, 16499),
    ("sr_pe_pool_demo", 16500, 16999),
    ("evpn_vpls_pool_demo", 10000, 10999),
]


def _rows(ws) -> list[tuple]:
    return [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]


def _write(ws, header: list[str], rows: list[tuple]) -> None:
    ws.delete_rows(1, ws.max_row)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))


def build(src: Path, dst: Path) -> None:
    wb = openpyxl.load_workbook(src)

    # Devices: rr + core, real SR OS models, no blank separator rows.
    ws = wb["Devices"]
    rows = _rows(ws)[1:]
    devices = [(n, role, site, MODEL_BY_ROLE[role], tenant) for n, role, site, _m, tenant in rows]
    _write(ws, ["DeviceName", "DeviceRole", "Site", "Model", "Tenant"], devices)

    # DistDevices: PE (pairs), one model.
    ws = wb["DistDevices"]
    rows = _rows(ws)[1:]
    dist = []
    for r in rows:
        name, role, site, _m, tenant = r[:5]
        name = ", ".join(p.strip() for p in str(name).split(","))
        dist.append((name, role, site, MODEL_BY_ROLE[role], tenant))
    _write(ws, ["DeviceName", "RoleName", "SiteName", "ModelName", "Tenant"], dist)

    # Role sheet: router roles plus the CE role.
    ws = wb["Role"]
    _write(ws, ["RoleName"], [("core",), ("pe",), ("rr",), ("switch",)])

    # Cables: ports are assigned automatically by the cable builder; keep as-is
    # but normalise to four columns.
    ws = wb["Cables"]
    rows = [(a, ia, b, ib) for a, ia, b, ib in _rows(ws)[1:]]
    _write(ws, ["Device_a", "Iface_a", "Device_b", "Iface_b"], rows)

    # Site: fictional addresses.
    ws = wb["Site"]
    sites = [r[0] for r in _rows(ws)[1:]]
    site_rows = [
        (s, f"Netwerkstraat {i + 1}", f"{1000 + i * 37:04d} AB", CITIES[i % len(CITIES)])
        for i, s in enumerate(sites)
    ]
    _write(ws, ["SiteName", "Address", "Postal Code", "City"], site_rows)

    ws = wb["PrefixPools"]
    _write(ws, ["PrefixPoolName", "PrefixPoolType", "Prefix"], PREFIX_POOLS)

    ws = wb["ResourcePools"]
    _write(ws, ["ResourcePoolName", "RangeStart", "RangeEnd"], RESOURCE_POOLS)

    # HalfOpenRings: trim trailing commas/spaces inside cells ("pe1.Site11, pe2.Site11, ").
    ws = wb["HalfOpenRings"]
    header = [c for c in _rows(ws)[0]]
    rings = []
    for r in _rows(ws)[1:]:
        cleaned = []
        for v in r:
            if isinstance(v, str):
                v = ", ".join(p.strip() for p in v.split(",") if p.strip())
            cleaned.append(v)
        rings.append(tuple(cleaned))
    _write(ws, header, rings)

    # CEs: switch model name, no ConnectedPE (attaches to the PE/pair at the site).
    ws = wb["CEs"]
    ces = [(n, site, "access-switch", "switch", None) for n, site, _m, _role, *_ in _rows(ws)[1:]]
    _write(ws, ["CEname", "SiteName", "ModelName", "CERole", "ConnectedPE"], ces)

    for name in list(wb.sheetnames):
        if name.startswith("report_"):
            del wb[name]

    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)


if __name__ == "__main__":
    build(Path(sys.argv[1]), Path(sys.argv[2]))
